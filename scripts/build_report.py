"""Build the submission report and diagrams from saved verification evidence.

Dependencies: python-docx, Pillow. Use the managed document runtime in Codex.
Technical figures are drawn from code; waveform captures use actual VCD events.
"""
from pathlib import Path
import csv
import json
import re
from datetime import date
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT/'docs'
BLUE = '#2E74B5'
INK = '#19324A'
PALE = '#E8EEF5'

def font(size, bold=False, mono=False):
    base = Path('C:/Windows/Fonts')
    name = 'consola.ttf' if mono else 'calibrib.ttf' if bold else 'calibri.ttf'
    if (base/name).exists():
        return ImageFont.truetype(str(base/name), size)
    try:
        return ImageFont.truetype('DejaVuSans.ttf',size)
    except OSError:
        return ImageFont.load_default()

def box(draw, rect, title, body='', fill=PALE):
    draw.rounded_rectangle(rect, radius=16, fill=fill, outline=BLUE, width=3)
    x0,y0,x1,y1=rect
    lines=[(title,font(29,True))]+[(line,font(24)) for line in body.split('\n') if line]
    start=(y0+y1)/2-len(lines)*19
    for index,(line,f) in enumerate(lines):
        draw.text(((x0+x1)/2,start+index*40),line,font=f,fill=INK,anchor='mt')

def arrow(draw, points, label='', label_at=None):
    draw.line(points,fill=INK,width=4)
    x,y=points[-1]; px,py=points[-2]
    if x==px:
        sign=1 if y>py else -1
        triangle=[(x,y),(x-9,y-sign*15),(x+9,y-sign*15)]
    else:
        sign=1 if x>px else -1
        triangle=[(x,y),(x-sign*15,y-9),(x-sign*15,y+9)]
    draw.polygon(triangle,fill=INK)
    if label:
        draw.text(label_at or points[0],label,font=font(22),fill=INK)

def diagrams():
    im=Image.new('RGB',(1600,1000),'white'); d=ImageDraw.Draw(im)
    d.text((50,35),'Project 2 | demonstration datapath',font=font(42,True),fill=INK)
    box(d,(50,150,360,280),'32-bit instruction','R / I / S encodings')
    box(d,(490,130,920,300),'Instruction decoder','rd, rs1, rs2, immediate\nop, src, rw, mr, mw, invalid')
    box(d,(60,440,425,600),'Register file','32 x 8 bits; x0 = 0')
    box(d,(550,450,830,590),'Operand MUX','rs2 / immediate[7:0]')
    box(d,(1010,430,1530,610),'Configurable 8-bit ALU','structural OR behavioral\nresult + C Z N V')
    box(d,(1010,740,1510,880),'Data memory','256 bytes; LB / SB')
    box(d,(525,740,840,880),'Writeback MUX','ALU / loaded byte')
    arrow(d,[(360,210),(490,210)],'instruction',(370,172))
    arrow(d,[(550,300),(550,350),(240,350),(240,440)],'register indices',(265,316))
    arrow(d,[(710,300),(710,450)],'imm / src',(725,352))
    arrow(d,[(920,210),(1260,210),(1260,430)],'ALU op',(1275,320))
    arrow(d,[(425,510),(550,510)],'rs2',(452,475))
    arrow(d,[(425,465),(465,465),(465,390),(1100,390),(1100,430)],'rs1',(880,352))
    arrow(d,[(830,525),(1010,525)],'operand B',(855,485))
    arrow(d,[(1330,610),(1330,740)],'address',(1350,695))
    arrow(d,[(1450,610),(1450,665),(685,665),(685,740)],'ALU result',(820,633))
    arrow(d,[(920,265),(970,265),(970,710),(1040,710),(1040,740)],'mr / mw',(885,680))
    arrow(d,[(1010,820),(840,820)],'read byte',(850,778))
    arrow(d,[(525,805),(245,805),(245,600)],'writeback / rw',(280,765))
    arrow(d,[(385,600),(385,930),(1190,930),(1190,880)],'rs2 store data; mw',(610,938))
    d.text((55,970),'clk / reset / enable gate state updates; invalid suppresses writes. ALU benchmark excludes the surrounding demo.',font=font(23),fill=INK)
    im.save(DOCS/'architecture.png')
    im=Image.new('RGB',(1600,650),'white'); d=ImageDraw.Draw(im)
    d.text((50,30),'Structural hierarchy and comparison boundary',font=font(42,True),fill=INK)
    labels=[('Half adder','XOR + AND'),('Full adder','2 half adders + OR'),('Ripple carry','WIDTH full adders'),('Adder/subtractor','B XOR sub; Cin=sub')]
    for index,(title,body) in enumerate(labels):
        x=40+index*390
        box(d,(x,160,x+350,305),title,body)
        if index<3: arrow(d,[(x+350,230),(x+390,230)])
    box(d,(100,430,740,575),'Structural ALU','arithmetic + gate logic + staged shifts + MUX tree')
    box(d,(860,430,1500,575),'Behavioral ALU','+, -, bitwise, shifts, signed compare + case')
    arrow(d,[(1385,305),(1385,360),(420,360),(420,430)])
    d.text((800,490),'=',font=font(46,True),fill=BLUE,anchor='mm')
    d.text((50,610),'Same input vectors, same flag contract, independent golden results; synthesis is run on each ALU alone.',font=font(25),fill=INK)
    im.save(DOCS/'structural_hierarchy.png')

def vcd_data():
    scope=[]; symbols={}; values={}; events={}; timestamp=0; in_header=True
    for raw in (ROOT/'results/demo.vcd').read_text().splitlines():
        line=raw.strip()
        if in_header:
            parts=line.split()
            if line.startswith('$scope'): scope.append(parts[2])
            elif line.startswith('$upscope'): scope.pop()
            elif line.startswith('$var') and scope==['tb_top']:
                symbols[parts[3]]=(parts[4],int(parts[2]))
            elif line.startswith('$enddefinitions'): in_header=False
            continue
        if not line or line.startswith('$'): continue
        if line.startswith('#'): timestamp=int(line[1:]); continue
        if line[0] in 'bB':
            value,symbol=line[1:].split()
        elif line[0] in '01xzXZ': value,symbol=line[0],line[1:]
        else: continue
        if symbol in symbols:
            name,width=symbols[symbol]
            # collapse zero-time delta activity into the final value at this timestamp
            events.setdefault(name,{})[timestamp]=value
    for name,changes in events.items():
        previous=None
        filtered={}
        for timestamp,value in sorted(changes.items()):
            if value!=previous: filtered[timestamp]=value
            previous=value
        events[name]=filtered
    return symbols,events

def waveform(path,start,end,title,notes):
    _,events=vcd_data()
    signals=['clk','enable','instruction','op','src','rw','mr','mw','invalid','a','b','result','wb','flags']
    W,H=1800,1050; left,right,top,spacing=170,1750,185,51
    im=Image.new('RGB',(W,H),'white'); d=ImageDraw.Draw(im)
    d.text((45,25),title,font=font(37,True),fill=INK)
    d.text((45,82),'Actual Icarus VCD | values after each timestamp settles | zero-delay RTL, not measured propagation',font=font(25),fill=INK)
    def x(t): return left+(t-start)/(end-start)*(right-left)
    tick=10 if end-start<100 else 20
    for t in range(start,end+1,tick):
        d.line([(x(t),top-20),(x(t),top+len(signals)*spacing)],fill='#DDE4EC',width=1)
        d.text((x(t),top-47),str(t),font=font(20),fill=INK,anchor='mt')
    d.text((right,top-78),'ns',font=font(21),fill=INK)
    for row,name in enumerate(signals):
        y=top+row*spacing
        d.text((left-18,y+15),name if name!='flags' else 'CZNV',font=font(23,mono=True),fill=INK,anchor='rm')
        ev=events.get(name,{})
        prior=[t for t in ev if t<=start*1000]
        current=ev[max(prior)] if prior else 'x'
        transitions=[(start,current)]+[(t/1000,v) for t,v in sorted(ev.items()) if start*1000<t<end*1000]+[(end,None)]
        for index,(t,value) in enumerate(transitions[:-1]):
            t2=transitions[index+1][0]; xa,xb=x(t),x(t2)
            if name in ['clk','enable','src','rw','mr','mw','invalid']:
                yy=y+4 if value=='1' else y+30
                d.line([(xa,yy),(xb,yy)],fill=BLUE,width=3)
                if index: d.line([(xa,y+4),(xa,y+30)],fill=BLUE,width=2)
            else:
                d.polygon([(xa,y+17),(xa+4,y+3),(xb-4,y+3),(xb,y+17),(xb-4,y+31),(xa+4,y+31)],outline=BLUE)
                try: label=f'{int(value,2):04b}' if name=='flags' else f'{int(value,2):08X}' if name=='instruction' else f'{int(value,2):02X}'
                except ValueError: label='X'
                if xb-xa > len(label)*13+5:
                    d.text(((xa+xb)/2,y+15),label,font=font(21,mono=True),fill=INK,anchor='mm')
    for i,note in enumerate(notes):
        d.text((45,925+i*35),note,font=font(24,True),fill=INK)
    im.save(path)

def style_doc(doc):
    # compact_reference_guide exact base tokens. Named overrides: title 28pt,
    # figure captions 9pt, dense trace/table cells 9pt with single line spacing.
    section=doc.sections[0]
    section.page_width=Inches(8.5); section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Inches(1)
    section.header_distance=section.footer_distance=Inches(.492)
    for name,size,color,before,after in [('Normal',11,'19324A',0,6),('Title',28,'19324A',0,12),
            ('Subtitle',13,'19324A',0,8),('Heading 1',16,'2E74B5',18,10),
            ('Heading 2',13,'2E74B5',14,7),('Heading 3',12,'1F4D78',10,5),('Caption',9,'19324A',4,4)]:
        s=doc.styles[name]; s.font.name='Calibri'; s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(color)
        s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after)
        s.paragraph_format.line_spacing=1.25
    for sname in ('List Bullet','List Number'):
        s=doc.styles[sname]; s.font.name='Calibri'; s.font.size=Pt(11)
        s.paragraph_format.space_after=Pt(4); s.paragraph_format.line_spacing=1.25
        s.paragraph_format.left_indent=Inches(.375); s.paragraph_format.first_line_indent=Inches(-.188)
    for level in doc.part.numbering_part.element.findall('.//'+qn('w:lvl')):
        ppr=level.find(qn('w:pPr'))
        if ppr is None: ppr=OxmlElement('w:pPr'); level.append(ppr)
        ind=ppr.find(qn('w:ind'))
        if ind is None: ind=OxmlElement('w:ind'); ppr.append(ind)
        ind.set(qn('w:left'),'540'); ind.set(qn('w:hanging'),'270')
        tabs=ppr.find(qn('w:tabs'))
        if tabs is None: tabs=OxmlElement('w:tabs'); ppr.append(tabs)
        tab=OxmlElement('w:tab'); tab.set(qn('w:val'),'num'); tab.set(qn('w:pos'),'540'); tabs.append(tab)
    header=section.header.paragraphs[0]
    header.text='COMPUTER ARCHITECTURE & ORGANIZATION  |  PROJECT 2'
    header.runs[0].font.size=Pt(8); header.runs[0].font.color.rgb=RGBColor.from_string('607080')
    footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Project 2  |  ')
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    for run in footer.runs: run.font.size=Pt(8)

def table(doc,headers,rows,widths=None):
    widths=widths or [9360//len(headers)]*len(headers)
    widths[-1]=9360-sum(widths[:-1])
    t=doc.add_table(rows=1,cols=len(headers)); t.autofit=False
    t.style='Table Grid'
    for c,text in zip(t.rows[0].cells,headers): c.text=str(text)
    for row in rows:
        for cell,text in zip(t.add_row().cells,row): cell.text=str(text)
    pr=t._tbl.tblPr
    tw=pr.find(qn('w:tblW')); tw.set(qn('w:type'),'dxa'); tw.set(qn('w:w'),'9360')
    ti=OxmlElement('w:tblInd'); ti.set(qn('w:w'),'120'); ti.set(qn('w:type'),'dxa'); pr.append(ti)
    margins=OxmlElement('w:tblCellMar')
    for side,width in [('top',80),('bottom',80),('start',120),('end',120)]:
        el=OxmlElement('w:'+side); el.set(qn('w:w'),str(width)); el.set(qn('w:type'),'dxa'); margins.append(el)
    pr.append(margins)
    for col,width in zip(t._tbl.tblGrid.gridCol_lst,widths): col.set(qn('w:w'),str(width))
    for index,row in enumerate(t.rows):
        for cell,width in zip(row.cells,widths):
            cell.width=Inches(width/1440)
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.0
                for run in p.runs: run.font.name='Calibri'; run.font.size=Pt(9); run.bold=(index==0)
            if index==0:
                shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E8EEF5'); cell._tc.get_or_add_tcPr().append(shade)
        # Keep individual rows intact while allowing long tables to span pages.
        cant=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(cant)
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    return t

def para(doc,text): return doc.add_paragraph(text)
def page(doc,title): doc.add_page_break(); doc.add_heading(title,1)
def picture(doc,name,caption,width=6.5):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(DOCS/name),width=Inches(width))
    doc.add_paragraph(caption,'Caption')

def main():
    verification=json.loads((ROOT/'results/verification_summary.json').read_text())
    if verification['mode']!='full' or not verification['passed']:
        raise SystemExit('Run a passing full verification suite before building the report.')
    generic=json.loads((ROOT/'results/yosys/comparison.json').read_text()) if (ROOT/'results/yosys/comparison.json').exists() else []
    listing=json.loads((ROOT/'vectors/program_listing.json').read_text())
    trace=list(csv.DictReader((ROOT/'results/execution_trace.csv').open()))
    diagrams()
    waveform(DOCS/'waveform_arithmetic.png',10,90,'Arithmetic: carry, signed overflow, and negative subtraction',[
        'Step 2, 30-40 ns: FF + 01 = 00; CZNV = 1100 (carry and zero).',
        'Step 4, 50-60 ns: 7F + 01 = 80; CZNV = 0011 (negative and signed overflow).',
        'Step 7, 80-90 ns: 05 - 0A = FB; CZNV = 0010 (negative, borrow).'])
    waveform(DOCS/'waveform_memory.png',210,350,'Memory and control: store, load, invalid instructions, and stalls',[
        'Step 20: SB writes FB at address 03; step 21: LB returns FB through the writeback MUX.',
        'Steps 27-30: invalid=1 suppresses writes; steps 31-32: enable=0 preserves state.',
        'Step 33: LB returns FB again, confirming the stalled store did not overwrite memory.'])
    doc=Document(); style_doc(doc)
    doc.core_properties.title='Structural vs. Behavioral ALU & Arithmetic Architecture Suite'
    doc.core_properties.subject='CAO Project 2 submission report'
    doc.core_properties.author='[Student name]'
    doc.add_paragraph('PROJECT 2','Subtitle')
    doc.add_paragraph('Structural vs. Behavioral ALU & Arithmetic Architecture Suite','Title')
    para(doc,'Computer Architecture and Organization | Software-only Verilog project')
    para(doc,'Student: [Name]    ID: [Student ID]\nSection / instructor: [Course details]')
    para(doc,f'Evidence generated: {date.today().isoformat()}')
    doc.add_heading('Objective and outcome',1)
    para(doc,'Implement the same arithmetic architecture using gate-level structural modeling and behavioral RTL, verify both against independent expectations, and compare their synthesized implementations under identical conditions.')
    para(doc,'The 8-bit suite passes all 524,288 operand/operation combinations. Smaller-width exhaustive tests, 16-bit boundary/random tests, decoder vectors, and a 37-step register/MUX/memory demonstration also pass. Both generic synthesized ALUs pass directed netlist tests.')
    table(doc,['Evidence','Saved outcome'],[
        ('8-bit ALU','524,288 vectors; zero failures'),('Decoder','4,224 vectors; zero failures'),
        ('Integration demo','37 steps; final registers/memory and reset checked'),
        ('Generic cells','Structural 352; behavioral 309 (Yosys no-ABC flow)'),
        ('Vivado FPGA timing/resources','Pending execution on a Vivado installation')],[2700,6660])
    para(doc,'The original CAO-PROJECT-2026.docx defines the project and shared 100-point rubric. The additional decoder and bus demonstration provides evidence for rubric items beyond the isolated ALU comparison.')
    page(doc,'1. Architecture and module hierarchy')
    picture(doc,'architecture.png','Figure 1. Decoder, registers, MUXes, configurable ALU and byte memory. Flags use CZNV order.',6.0)
    picture(doc,'structural_hierarchy.png','Figure 2. Structural arithmetic hierarchy and the equal-interface behavioral implementation.',6.0)
    para(doc,'Structural full adders contain two half adders and an OR gate. A WIDTH-stage carry chain adds A and B XOR sub, with sub as the initial carry. Carry into and out of the sign bit are XORed to detect signed overflow. Logic operations use gates; logical shifts use fixed shift stages selected by gate MUXes. Behavioral RTL uses complete combinational blocks and standard operators.')
    page(doc,'2. Interface and arithmetic semantics')
    para(doc,'Both ALUs accept a[WIDTH-1:0], b[WIDTH-1:0] and op[2:0], and produce result[WIDTH-1:0], carry, zero, negative and overflow. The wrapper parameter STRUCTURAL chooses the implementation at elaboration. WIDTH defaults to 8; supported widths are powers of two of at least 2, verified at 4, 8 and 16 bits.')
    table(doc,['op','Operation','Result / interpretation'],[
        (0,'ADD','A + B, modulo 2^WIDTH'),(1,'SUB','A - B, modulo 2^WIDTH'),
        (2,'AND','Bitwise conjunction'),(3,'OR','Bitwise disjunction'),(4,'XOR','Bitwise exclusive OR'),
        (5,'SLL','Zero-fill left shift; discard high bits'),(6,'SRL','Zero-fill right shift'),
        (7,'SLT','Signed two\'s-complement comparison; result 0 or 1')],[600,1700,7060])
    doc.add_heading('Status flags',2)
    table(doc,['Flag','Definition'],[
        ('C','ADD: unsigned carry-out. SUB: 1 when A >= B (no borrow). Other operations: 0.'),
        ('Z','1 exactly when the result is zero, for every operation.'),
        ('N','Most significant result bit, for every operation.'),
        ('V','Signed overflow for ADD/SUB; cleared for logic, shifts and SLT.')],[700,8660])
    para(doc,'The shift amount is the low log2(WIDTH) bits of B. For the 8-bit ALU, a shift operand of 8 selects a shift of zero. The demonstration uses this same rule, including shift-immediate instructions. Flags describe the ALU result; they are combinational outputs rather than saved processor status.')
    table(doc,['Calculation','Result','CZNV','Reason'],[
        ('FF + 01','00','1100','Unsigned carry and zero'),('7F + 01','80','0011','Signed positive overflow'),
        ('05 - 0A','FB','0010','Negative result with borrow'),('0A - 0A','00','1100','Zero; no borrow')],[2200,1100,1000,5060])
    page(doc,'3. Decoder and demonstration execution')
    para(doc,'A small execution harness supplies evidence for the shared decoder and MUX rubric. It accepts a 32-bit instruction from the testbench, reads two 8-bit registers, selects register/immediate operand B, computes the ALU output and commits writes on an enabled rising edge. There are 32 registers and 256 data-memory bytes; x0 remains zero.')
    table(doc,['Encoding family','Supported instructions'],[
        ('R-type','ADD, SUB, AND, OR, XOR, SLL, SRL, SLT'),
        ('I-type arithmetic','ADDI, ANDI, ORI, XORI, SLLI, SRLI, SLTI'),
        ('I-type load','LB'),('S-type store','SB')],[2400,6960])
    para(doc,'The decoder extracts rd, rs1, rs2, funct3 and funct7 and validates opcode/function combinations. I/S immediates are sign-extended to 32 bits, then narrowed at the operand MUX. Byte addresses wrap modulo 256. LB writes its eight-bit byte unchanged into the eight-bit register file. These are educational datapath semantics using RISC-V encodings, not a complete RV32 implementation.')
    para(doc,'Unsupported instructions assert invalid and clear all write controls. enable=0 blocks commits without changing the combinational decode. Synchronous reset clears registers and memory and overrides enable. Nonblocking assignments are used for state; complete combinational assignments prevent latches.')
    doc.add_heading('Scope of the demonstration',2)
    para(doc,'The testbench supplies instructions directly. There is no program counter, instruction fetch, branching, pipeline, RV32 compliance claim or physical board interface. Each valid enabled instruction commits on one edge. This edge count is a property of the demonstration harness, not a performance comparison between the ALUs.')
    para(doc,'The file-driven program deliberately covers carry, signed overflow, borrowing, logic zeros, both shifts, signed comparisons, memory writes/reads, negative offsets, truncation, x0 writes, illegal instructions and disabled register/memory writes. Final state is compared against independently generated software expectations.')
    page(doc,'4. Automated verification and saved results')
    para(doc,f'Tool: {verification["tool_version"]}. Testbench-only assertion extensions are compiled with -g2012; the datapath sources remain synthesizable Verilog. Every run produces a process exit status, logs and source SHA-256 hashes in verification_summary.json.')
    counts={'tb_units':'16 truth-table checks','tb_alu_4':'2,048 exhaustive vectors','tb_alu_8':'524,288 exhaustive vectors','tb_alu_16':'16,896 boundary/random vectors','tb_decoder':'4,224 file vectors','tb_top':'37 instructions + final-state/reset checks'}
    table(doc,['Test','Coverage','Result','Runtime (s)'],[(r['name'],counts[r['name']],'PASS' if r['passed'] else 'FAIL',r['simulation_seconds']) for r in verification['tests']],[1800,4700,900,1960])
    para(doc,'For widths 4 and 8, all A/B pairs are checked for all eight operations. ADD/SUB also check both standalone adder/subtractor modules. The reference converts operands to signed integers and checks signed mathematical range for overflow; it does not copy DUT carry-chain equations. Sixteen-bit testing uses an 8-by-8 boundary grid and 2,048 seeded random pairs per operation.')
    para(doc,'The decoder suite traverses all funct7/funct3 combinations in four opcode families plus all opcode values. Integration uses external .mem expectations generated by Python. Case-inequality comparisons detect unknown/high-impedance outputs as mismatches. The suite checks all 32 final register bytes and 256 final memory bytes in both styles, then checks reset clears them.')
    para(doc,'Console logging uses $display; mismatches issue $error and/or $fatal and stop the failing run. The Python runner returns a nonzero status if compilation or verification fails. External vectors use $readmemh. Exhaustive tests avoid dumping large waveforms; the directed integration run saves demo.vcd.')
    para(doc,'Runtime is measured software simulation wall time on this machine. The 8-bit ALU run advances 524,288 ns using a 1 ns testbench settle interval. This is an artificial test schedule, not a measured circuit propagation delay or hardware cycle count.')
    page(doc,'5. Annotated waveform evidence')
    picture(doc,'waveform_arithmetic.png','Figure 3. Arithmetic transitions from the saved VCD, annotated with carry, signed overflow and negative-result examples.')
    para(doc,'These captures are generated from actual VCD signal changes and show settled values at each timestamp. The capture highlights the document\'s carry/zero and negative-result examples, plus signed overflow. Testbench time intervals are chosen for readability and do not model physical propagation delay.')
    page(doc,'5. Waveform evidence: memory and control')
    picture(doc,'waveform_memory.png','Figure 4. Actual bus/control traces around stores, loads, invalid instructions and stalls. Short bus intervals may omit labels; the CSV contains exact values.')
    para(doc,'These captures are generated from VCD signal changes, collapsed to the settled value at each timestamp. They visualize functional activity and contain no physical gate-delay model. Open docs/demo.gtkw in GTKWave for the full interactive trace. Trace rows are sampled before the committing edge, after the input settle interval.')
    for half in (0,1):
        page(doc,f'6. Execution trace ({half+1}/2)')
        entries=list(zip(trace,listing))[half*19:(half+1)*19]
        rows=[]
        for actual,expected in entries:
            rows.append((actual['step'],expected['assembly'],actual['a'],actual['b'],actual['result'],actual['writeback'],actual['CZNV']))
        table(doc,['#','Instruction / scenario','A','B','Y','WB','CZNV'],rows,[650,5010,650,650,650,650,1100])
        para(doc,'All values in A/B/Y/WB are hexadecimal. CZNV is binary. Y is the ALU result and WB is the selected writeback bus; WB does not imply that a register write is enabled. Stores use the ALU output as an address. Unsupported or stalled rows do not commit state changes.')
        para(doc,'Source: results/execution_trace.csv is emitted by the actual Verilog simulation. Assembly descriptions are joined from vectors/program_listing.json, generated by the independent software model. The complete CSV additionally records time, instruction encoding, register indices, immediate, ALU selection and write controls.')
    page(doc,'7. Synthesis comparison and timing method')
    para(doc,'The saved generic comparison uses identical flattened Yosys synthesis with synth -noabc. This maps operations to generic Boolean cells before ABC optimization. It checks connectivity, rejects inferred latches and produces a functional netlist for each ALU. Both netlists pass 24 directed vectors using Icarus.')
    if generic:
        table(doc,['Implementation','Generic cells','Netlist test'],[(r['style'],r['generic_cells'],'PASS' if r['post_synthesis_passed'] else 'FAIL') for r in generic],[3500,2500,3360])
    para(doc,'In this flow the behavioral implementation uses 43 fewer generic logic cells, about 12.2% fewer than the structural result. Hierarchy metadata ($scopeinfo) is excluded. This is a measured outcome for the chosen tool and settings, not a claim of minimum gate count or universally better hardware. Generic cell primitives have different functions and fan-in; they are not standardized NAND-equivalent area units.')
    para(doc,'Behavioral operators expose arithmetic/comparison intent to synthesis. The structural circuit explicitly expresses a ripple chain, a signed comparison using subtract/overflow, and a MUX network; synthesis can still simplify this hierarchy. Technology mapping and optimization can change either result substantially.')
    doc.add_heading('Vivado method and current status',2)
    para(doc,'Vivado was unavailable for execution. The supplied Tcl flow targets Artix-7 xc7a35tcpg236-1 by default, synthesizes each ALU independently out of context, and applies the same 10 ns maximum input-to-output delay constraint. It reports synthesis resources, routes internal logic and exports utilization, critical-path timing, checkpoints, functional/timing netlists and SDF. No board IO buffers are included.')
    para(doc,'The flow also runs XSim post-synthesis functional tests using vendor primitive libraries. SDF export prepares optional timing simulation; it does not prove that timing simulation has been executed. Routed static timing will supply the device-specific propagation comparison. Existing generic counts must not be copied into FPGA LUT/CARRY4 columns.')
    para(doc,'The ALUs are combinational: an execution-cycle speedup is not defined here. Test-vector count, simulated time and wall runtime are reported separately. A timing-aware comparison requires the same target, constraints, tool version and optimization settings for both implementations.')
    page(doc,'8. Rubric coverage and reproduction')
    table(doc,['Category','Points','Evidence'],[
        ('Functional correctness','40','Exhaustive ALU/add-sub results and flags; R/I/S decoder; register/MUX/memory integration'),
        ('Automated verification','20','Self-checking tb_top, .mem vectors, structured loops, independent expectations, failure exit codes'),
        ('Code quality / synthesis','20','Parameterized hierarchy, complete combinational blocks, nonblocking state, Yosys latch/connectivity checks'),
        ('Simulation / waveforms','10','Saved VCD/GTKWave config, annotated arithmetic and bus/control captures'),
        ('Documentation / diagram','10','Architecture diagrams, interface contract, full trace, this report and demo guide')],[2500,800,6060])
    doc.add_heading('Reproduce the evidence',2)
    for text in ['Run python scripts/run_tests.py for the full portable suite.',
                 'Optionally run python scripts/run_synthesis.py with Yosys/YoWASP installed for generic netlists and cell counts.',
                 'In Vivado source scripts/create_vivado_project.tcl, then run behavioral simulation with tb_top as the top.',
                 'Run vivado -mode batch -source scripts/vivado_compare.tcl for isolated FPGA measurements and XSim netlist checks.',
                 'Fill in the report identity fields and attach genuine Vivado outputs if the course requires them.']:
        doc.add_paragraph(text,'List Number')
    doc.add_heading('Sources and conclusion',2)
    para(doc,'Project authority: CAO-PROJECT-2026.docx, Detailed Project Description 2, Testing Methodology and Project Grading Rubric. Tool references: Icarus installation documentation (github.com/steveicarus/iverilog); AMD UG835 write_sdf and UG900 Generating a Timing Netlist (docs.amd.com). Exact links are in README.md.')
    para(doc,'Both coding styles satisfy the same tested arithmetic contract. The delivered generic synthesis demonstrates a resource difference under stated settings. FPGA timing and resource conclusions remain pending the supplied Vivado workflow; the functional results and saved generic synthesis evidence have been executed and verified.')
    output=DOCS/'Project_2_Report.docx'; doc.save(output)
    # Deterministic structural audit of the preset's fixed table geometry.
    for t in doc.tables:
        grid=[int(c.get(qn('w:w'))) for c in t._tbl.tblGrid.gridCol_lst]
        assert sum(grid)==9360
        for row in t.rows:
            assert [int(c._tc.get_or_add_tcPr().find(qn('w:tcW')).get(qn('w:w'))) for c in row.cells]==grid
    print(f'Created {output}; {len(doc.tables)} tables passed geometry audit.')

if __name__=='__main__': main()
